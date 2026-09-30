#!/usr/bin/env python3
"""Render the analysis-focused response source to an editable Word draft.

Requires python-docx. The Markdown is the version-controlled prose source.
Use the documents skill renderer separately for page-by-page layout QA.
"""
from pathlib import Path
import argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build(source: Path, output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    # The bundled default template can carry a colored Title border.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name, size in [('Normal', 11), ('Title', 19), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.06
        style.paragraph_format.widow_control = True
        if name != 'Normal':
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)

    text = source.read_text(encoding='utf-8')
    for block in text.strip().split('\n\n'):
        block = block.strip()
        if block.startswith('# '):
            doc.add_paragraph(block[2:], 'Title')
        elif block.startswith('## '):
            title = block[3:]
            p = doc.add_paragraph(title, 'Heading 1')
        elif block.startswith('### '):
            title = block[4:]
            p = doc.add_paragraph(title, 'Heading 2')
            notes = {
                'Short intertrial intervals and model separability':
                    'Timing clarification is supported. The response does not yet resolve the first-level rank flags; see the companion author notes before finalizing.',
                'Number of regressors and alternative models':
                    'Open diagnostic: the production summary flags 50 of 94 runs per model family. Identify zero columns and focal-contrast estimability before claiming the model is unaffected. No new imaging fit is implied.',
                'Simple effects underlying the DMN interaction':
                    'The tables and four condition bars exist. The reviewer also requested separate age-specific brain renderings; those additional displays have not been verified as completed.',
                'A potentially influential older participant':
                    'Author decision: retain this explanation or authorize a separate complete-sample outlier-deweighting sensitivity analysis. The existing historical ROI audit does not answer corrected whole-brain influence.',
            }
            if title in notes:
                doc.add_comment(p.runs, text=notes[title], author='Author review')
        elif block.startswith('> '):
            p = doc.add_paragraph(block[2:])
            p.paragraph_format.left_indent = Inches(.2)
            p.paragraph_format.right_indent = Inches(.12)
        elif block.startswith('Proposed '):
            p = doc.add_paragraph()
            p.add_run(block).bold = True
            p.paragraph_format.keep_with_next = True
        elif block.startswith('Concern summary:'):
            p = doc.add_paragraph()
            p.add_run(block).italic = True
            p.paragraph_format.keep_with_next = True
        else:
            doc.add_paragraph(block)

    p = section.footer.paragraphs[0]
    p.alignment = 2
    run = p.add_run()
    run.font.size = Pt(9)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run._r.addnext(fld)
    doc.core_properties.title = 'Responses to analytical and results concerns'
    doc.core_properties.author = 'SRNDNA Ultimatum Game authors'
    doc.core_properties.subject = 'First author review pass with proposed manuscript passages'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == '__main__':
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=root / 'docs/revision/analysis_responses.md')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)
